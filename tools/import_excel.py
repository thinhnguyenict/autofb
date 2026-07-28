"""One-time importer from a reviewable Excel workbook into AutoFB posts."""
from __future__ import annotations

import argparse
import os
from datetime import UTC, date, datetime
from pathlib import Path
from typing import Any, Iterable

from autofb.web.database import Database
from autofb.web.service import AutoFBService, ServiceError

REQUIRED_COLUMNS = {"facebook_page_id", "body"}


def _text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value).strip()


def read_workbook(path: str | Path) -> list[dict[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        iterator = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip().lower().replace(" ", "_") for value in next(iterator, ())]
        if not REQUIRED_COLUMNS.issubset(headers):
            missing = ", ".join(sorted(REQUIRED_COLUMNS - set(headers)))
            raise ValueError(f"Workbook is missing required columns: {missing}")
        rows = []
        for values in iterator:
            row = {header: _text(value) for header, value in zip(headers, values) if header}
            if any(row.values()):
                rows.append(row)
        return rows
    finally:
        workbook.close()


def import_rows(
    database: Database,
    actor_email: str,
    workspace_name: str,
    rows: Iterable[dict[str, str]],
    *,
    dry_run: bool = False,
) -> dict[str, int]:
    database.initialize()
    service = AutoFBService(database)
    with database.connect() as connection:
        actor = connection.execute("SELECT id FROM users WHERE email = ?", (actor_email.strip().lower(),)).fetchone()
        if actor is None:
            raise ServiceError("Import actor is not registered")
        workspace = connection.execute(
            """SELECT workspaces.id, workspace_members.role FROM workspaces
               JOIN workspace_members ON workspace_members.workspace_id = workspaces.id
               WHERE workspaces.name = ? AND workspace_members.user_id = ?""",
            (workspace_name.strip(), actor["id"]),
        ).fetchone()
        if workspace is None:
            raise ServiceError("Workspace is missing or the import actor is not a member")
    pages = {page["facebook_page_id"]: page for page in service.list_facebook_pages(actor["id"], workspace["id"])}
    prepared = []
    for number, raw in enumerate(rows, start=2):
        page_id = _text(raw.get("facebook_page_id"))
        body = _text(raw.get("body"))
        scheduled_at = _text(raw.get("scheduled_at"))
        timezone = _text(raw.get("timezone")) or "UTC"
        if not page_id or page_id not in pages:
            raise ServiceError(f"Row {number}: Facebook Page is not connected to the workspace")
        if not body:
            raise ServiceError(f"Row {number}: post body is required")
        if scheduled_at:
            try:
                parsed = datetime.fromisoformat(scheduled_at.replace("Z", "+00:00"))
            except ValueError as exc:
                raise ServiceError(f"Row {number}: scheduled_at must be ISO-8601") from exc
            if parsed.tzinfo is None:
                raise ServiceError(f"Row {number}: scheduled_at must include a timezone")
            if parsed <= datetime.now(UTC):
                raise ServiceError(f"Row {number}: scheduled_at must be in the future")
            if workspace["role"] == "editor":
                raise ServiceError("Editor imports cannot schedule posts until each draft is approved")
        prepared.append((pages[page_id]["id"], body, scheduled_at, timezone))

    if dry_run:
        return {"validated": len(prepared), "created": 0, "scheduled": 0}
    created = scheduled = 0
    for page_id, body, scheduled_at, timezone in prepared:
        post = service.create_post(actor["id"], workspace["id"], page_id, body)
        created += 1
        if scheduled_at:
            service.schedule_post(actor["id"], workspace["id"], post["id"], scheduled_at, timezone)
            scheduled += 1
    return {"validated": len(prepared), "created": created, "scheduled": scheduled}


def parser() -> argparse.ArgumentParser:
    argument_parser = argparse.ArgumentParser(description="Import AutoFB drafts/schedules from an Excel workbook once.")
    argument_parser.add_argument("workbook", help="XLSX file with facebook_page_id and body columns.")
    argument_parser.add_argument("--database", default=os.environ.get("AUTOFB_DATABASE_PATH", "autofb.db"))
    argument_parser.add_argument("--actor-email", default=os.environ.get("AUTOFB_IMPORT_ACTOR_EMAIL", ""))
    argument_parser.add_argument("--workspace", default=os.environ.get("AUTOFB_IMPORT_WORKSPACE", ""))
    argument_parser.add_argument("--dry-run", action="store_true")
    return argument_parser


def main() -> int:
    args = parser().parse_args()
    if not args.actor_email or not args.workspace:
        raise SystemExit("--actor-email and --workspace (or AUTOFB_IMPORT_ACTOR_EMAIL/AUTOFB_IMPORT_WORKSPACE) are required")
    result = import_rows(Database(args.database), args.actor_email, args.workspace, read_workbook(args.workbook), dry_run=args.dry_run)
    print(f"Excel import complete: validated={result['validated']} created={result['created']} scheduled={result['scheduled']}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
